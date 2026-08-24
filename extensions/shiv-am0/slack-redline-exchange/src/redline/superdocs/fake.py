from __future__ import annotations

import base64
import html as html_mod
import io
import json
import re
from typing import Any

from .models import (
    AWAITING_KIND_CONTINUE,
    OPEN_MODE_BACKGROUND,
    OPEN_MODE_REPLACE,
    OPEN_MODES,
    STATUS_AWAITING_APPROVAL,
    STATUS_COMPLETED,
    STATUS_IN_PROGRESS,
    STATUS_PENDING,
)
from .transport import SuperDocsAPIError


class FakeSession:
    def __init__(self, session_id: str, filename: str):
        self.id = session_id
        self.filename = filename
        self.version = 0
        self.documents: dict[str, FakeDocument] = {}
        self.focused_document_id: str | None = None

    @property
    def focused(self) -> FakeDocument:
        if self.focused_document_id is None:
            raise SuperDocsAPIError(
                409, "no_document", f"session {self.id!r} has no open document"
            )
        return self.documents[self.focused_document_id]

    @property
    def blocks(self) -> list[dict[str, str]]:
        return self.focused.blocks

    @blocks.setter
    def blocks(self, value: list[dict[str, str]]) -> None:
        self.focused.blocks = value

    def document_or_focused(self, document_id: str | None) -> FakeDocument:
        if document_id is None:
            return self.focused
        if document_id not in self.documents:
            raise SuperDocsAPIError(
                404,
                "document_not_found",
                f"unknown document {document_id!r} in session {self.id!r}; "
                f"open documents are {sorted(self.documents)}",
            )
        return self.documents[document_id]


class FakeDocument:
    def __init__(self, document_id: str, title: str):
        self.id = document_id
        self.title = title
        self.blocks: list[dict[str, str]] = []


class FakeJob:
    def __init__(
        self,
        job_id: str,
        session_id: str,
        message: str,
        approval_mode: str,
        document_id: str | None = None,
    ):
        self.id = job_id
        self.session_id = session_id
        self.message = message
        self.approval_mode = approval_mode
        self.document_id = document_id
        self.status = STATUS_PENDING
        self.stage = 0
        self.round = 0
        self.proposals: list[dict[str, Any]] = []
        self.decisions_log: list[dict[str, Any]] = []
        self.result: dict[str, Any] | None = None
        self.error: str | None = None
        self.awaiting_kind: str | None = None


class FakeSuperDocs:
    def __init__(self, ops_limit: int = 100):
        self.sessions: dict[str, FakeSession] = {}
        self.jobs: dict[str, FakeJob] = {}
        self.records: list[tuple[str, str, dict[str, Any]]] = []
        self.ops_used = 0
        self.ops_limit = ops_limit
        self._session_counter = 0
        self._job_counter = 0
        self._chunk_counter = 0
        self._document_counter = 0

    def _record(self, method: str, path: str, payload: dict[str, Any]) -> None:
        self.records.append((method, path, payload))

    def _next_id(self, prefix: str) -> str:
        self._job_counter += 1
        return f"{prefix}_{self._job_counter:04d}"

    def _new_chunk(self, tag: str, text: str) -> dict[str, str]:
        self._chunk_counter += 1
        return {"chunk_id": f"c{self._chunk_counter:04d}", "tag": tag, "text": text}

    async def post(self, path: str, payload: dict[str, Any]) -> dict[str, Any]:
        self._record("POST", path, payload)
        if path == "/v1/documents/upload-base64":
            return self._upload(payload)
        if path == "/v1/chat/async":
            return self._chat_async(payload)
        if re.fullmatch(r"/v1/chat/[^/]+/approve", path):
            session_id = path.split("/")[3]
            return self._approve(session_id, payload)
        if re.fullmatch(r"/v1/chat/[^/]+/continue", path):
            session_id = path.split("/")[3]
            return self._continue(session_id, payload)
        if re.fullmatch(r"/v1/jobs/[^/]+/cancel", path):
            return self._cancel(path.split("/")[3])
        if path == "/v1/documents/export":
            return self._export(payload)
        focus = re.fullmatch(r"/v1/sessions/([^/]+)/documents/([^/]+)/focus", path)
        if focus:
            return self._focus_document(focus.group(1), focus.group(2))
        raise SuperDocsAPIError(404, "not_found", f"no fake route for {path}")

    async def get(self, path: str) -> dict[str, Any]:
        self._record("GET", path, {})
        if path.startswith("/v1/jobs/"):
            job_id = path.split("/")[3]
            return self._get_job(job_id)
        match = re.fullmatch(r"/v1/sessions/([^/]+)/jobs", path)
        if match:
            return self._list_session_jobs(match.group(1))
        documents = re.fullmatch(r"/v1/sessions/([^/]+)/documents", path)
        if documents:
            return self._list_documents(documents.group(1))
        raise SuperDocsAPIError(404, "not_found", f"no fake route for {path}")

    def _upload(self, payload: dict[str, Any]) -> dict[str, Any]:
        filename = payload["filename"]
        raw = base64.b64decode(payload["file_base64"])
        session_id = payload.get("session_id") or self._next_id("sess")
        open_mode = payload.get("open_mode") or OPEN_MODE_REPLACE
        if open_mode not in OPEN_MODES:
            raise SuperDocsAPIError(
                422, "validation_error", f"unknown open_mode {open_mode!r}"
            )
        session = self.sessions.get(session_id) or FakeSession(session_id, filename)
        blocks = [
            self._new_chunk(block["tag"], block["text"])
            for block in self._parse_file(filename, raw)
        ]

        if open_mode == OPEN_MODE_REPLACE and session.focused_document_id is not None:
            document = session.focused
            document.title = filename
            document.blocks = blocks
        else:
            self._document_counter += 1
            document = FakeDocument(f"doc_{self._document_counter:04d}", filename)
            document.blocks = blocks
            session.documents[document.id] = document
            if open_mode != OPEN_MODE_BACKGROUND or session.focused_document_id is None:
                session.focused_document_id = document.id

        session.filename = session.focused.title
        session.version += 1
        self.sessions[session_id] = session
        return {
            "session_id": session_id,
            "document_id": document.id,
            "chunks_count": len(document.blocks),
            "version_id": f"v{session.version}",
            "focused_document_id": session.focused_document_id,
        }

    def _list_documents(self, session_id: str) -> dict[str, Any]:
        session = self.sessions.get(session_id)
        if session is None:
            raise SuperDocsAPIError(
                404, "document_not_found", f"unknown session {session_id!r}"
            )
        return {
            "documents": [
                {
                    "document_id": doc.id,
                    "title": doc.title,
                    "is_focused": doc.id == session.focused_document_id,
                    "sections_count": len(doc.blocks),
                }
                for doc in session.documents.values()
            ],
            "focused_document_id": session.focused_document_id,
        }

    def _focus_document(self, session_id: str, document_id: str) -> dict[str, Any]:
        session = self.sessions.get(session_id)
        if session is None:
            raise SuperDocsAPIError(
                404, "document_not_found", f"unknown session {session_id!r}"
            )
        session.document_or_focused(document_id)
        session.focused_document_id = document_id
        return {"status": "ok", "focused_document_id": document_id}

    @staticmethod
    def _parse_file(filename: str, raw: bytes) -> list[dict[str, str]]:
        blocks: list[dict[str, str]] = []
        lower = filename.lower()
        if lower.endswith((".md", ".txt")):
            buffer: list[str] = []

            def flush() -> None:
                if buffer:
                    blocks.append({"tag": "p", "text": " ".join(buffer)})
                    buffer.clear()

            for line in raw.decode("utf-8").splitlines():
                stripped = line.strip()
                if not stripped:
                    flush()
                    continue
                if stripped.startswith("#"):
                    flush()
                    level = len(stripped) - len(stripped.lstrip("#"))
                    heading_text = stripped.lstrip("# ").strip()
                    blocks.append({"tag": f"h{min(level, 6)}", "text": heading_text})
                else:
                    buffer.append(stripped)
            flush()
        elif lower.endswith(".html") or lower.endswith(".htm"):
            text = raw.decode("utf-8")
            tag_pattern = r"<(?:p|h[1-6]|li)[^>]*>(.*?)</"
            for part in re.findall(tag_pattern, text, re.DOTALL | re.IGNORECASE):
                clean = re.sub(r"<[^>]+>", "", part).strip()
                if clean:
                    blocks.append({"tag": "p", "text": clean})
        elif lower.endswith(".docx"):
            import docx

            document = docx.Document(io.BytesIO(raw))
            for para in document.paragraphs:
                style = (para.style.name or "").lower()
                if para.text.strip():
                    tag = "h1"
                    for size in range(1, 7):
                        if f"heading {size}" in style:
                            tag = f"h{size}"
                            break
                    else:
                        tag = "p"
                    blocks.append({"tag": tag, "text": para.text.strip()})
        else:
            for chunk in raw.decode("utf-8", errors="replace").split("\n\n"):
                if chunk.strip():
                    blocks.append({"tag": "p", "text": chunk.strip()})
        return blocks or [{"tag": "p", "text": "(empty document)"}]

    def _chat_async(self, payload: dict[str, Any]) -> dict[str, Any]:
        session_id = payload.get("session_id")
        if not session_id or session_id not in self.sessions:
            raise SuperDocsAPIError(404, "document_not_found", f"unknown session {session_id!r}")
        if self.ops_used >= self.ops_limit:
            raise SuperDocsAPIError(
                429, "operation_limit_reached", "monthly operation limit reached"
            )
        message = payload.get("message", "")
        if not isinstance(message, str) or not message.strip():
            raise SuperDocsAPIError(422, "validation_error", "message must be non-empty text")
        self.ops_used += 1
        session = self.sessions[str(session_id)]
        target = session.document_or_focused(payload.get("document_id"))
        if payload.get("document_id"):
            session.focused_document_id = target.id
        job = FakeJob(
            job_id=self._next_id("job"),
            session_id=str(session_id),
            message=message.strip(),
            approval_mode=payload.get("approval_mode", "auto"),
            document_id=target.id,
        )
        self.jobs[job.id] = job
        return {
            "job_id": job.id,
            "session_id": session_id,
            "document_id": target.id,
            "status": STATUS_PENDING,
        }

    def _advance(self, job: FakeJob) -> None:
        if job.status == STATUS_PENDING:
            job.status = STATUS_IN_PROGRESS
            return
        if job.status == STATUS_IN_PROGRESS:
            session = self.sessions[job.session_id]
            document = session.document_or_focused(job.document_id)
            if job.approval_mode == "ask_every_time":
                job.proposals = self._draft_proposals(job.message, document)
                job.status = STATUS_AWAITING_APPROVAL
                job.awaiting_kind = None
            else:
                drafted = self._draft_proposals(job.message, document)
                self._apply(session, [p["change_id"] for p in drafted], drafted, document)
                job.status = STATUS_COMPLETED
                job.result = {
                    "response": f"Applied {len(drafted)} change(s).",
                    "document_changes": {
                        "updated_html": self.render_html(job.session_id, document.id),
                        "version_id": f"v{session.version}",
                        "changes_summary": f"{len(drafted)} change(s) applied",
                        "document_id": document.id,
                    },
                    "usage": {
                        "monthly_used": self.ops_used,
                        "monthly_remaining": self.ops_limit - self.ops_used,
                    },
                }

    def _get_job(self, job_id: str) -> dict[str, Any]:
        job = self.jobs.get(job_id)
        if job is None:
            raise SuperDocsAPIError(404, "job_not_found", f"unknown job {job_id!r}")
        if job.status in (STATUS_PENDING, STATUS_IN_PROGRESS):
            self._advance(job)
        meta: dict[str, Any] = {}
        if job.status == STATUS_AWAITING_APPROVAL:
            if job.awaiting_kind == AWAITING_KIND_CONTINUE:
                meta["awaiting_kind"] = AWAITING_KIND_CONTINUE
                meta["continue_prompt"] = {
                    "message": "Large edit paused; continue?",
                    "done": job.round * 10,
                    "total": 100,
                    "remaining": 100 - job.round * 10,
                }
            else:
                encoded = json.dumps(
                    {"type": "single_approval", "changes": job.proposals}
                )
                meta["awaiting_kind"] = "approval_review"
                meta["pending_changes"] = encoded
        payload: dict[str, Any] = {"job_id": job.id, "status": job.status, "metadata": meta}
        if job.result is not None:
            payload["result"] = job.result
        if job.error is not None:
            payload["error"] = job.error
        return payload

    def _approve(self, session_id: str, payload: dict[str, Any]) -> dict[str, Any]:
        if "approved" not in payload:
            raise SuperDocsAPIError(
                422, "validation_error", "top-level 'approved' field is required"
            )
        job_id = payload.get("job_id")
        job = self.jobs.get(str(job_id)) if job_id else None
        if job is None:
            raise SuperDocsAPIError(404, "job_not_found", f"unknown job {job_id!r}")
        if job.session_id != session_id:
            raise SuperDocsAPIError(
                409, "session_mismatch", "job does not belong to this session"
            )
        if job.status != STATUS_AWAITING_APPROVAL or job.awaiting_kind == AWAITING_KIND_CONTINUE:
            raise SuperDocsAPIError(409, "not_awaiting_approval", "job is not paused for review")
        changes = payload.get("changes") or [
            {"change_id": p["change_id"], "approved": bool(payload["approved"])}
            for p in job.proposals
        ]
        job.decisions_log.extend(changes)
        accepted_ids = [
            str(c["change_id"])
            for c in changes
            if bool(c.get("approved", payload["approved"]))
        ]
        feedback_items = [
            c for c in changes if not c.get("approved", True) and c.get("feedback")
        ]
        session = self.sessions[session_id]
        document = session.document_or_focused(job.document_id)
        valid_ids = {p["change_id"] for p in job.proposals}
        unknown = set(accepted_ids) - valid_ids
        if unknown:
            raise SuperDocsAPIError(
                404, "change_not_found", f"unknown change ids: {sorted(unknown)}"
            )
        applied = self._apply(session, accepted_ids, job.proposals, document)
        job.applied = getattr(job, "applied", []) + applied
        job.proposals = []
        if applied or accepted_ids:
            job.status = STATUS_COMPLETED
            job.awaiting_kind = None
            denied = len(changes) - len(accepted_ids)
            job.result = {
                "response": (
                    f"Applied {len(applied)} approved change(s)"
                    + (f"; discarded {denied} denied change(s)" if denied else "")
                    + "."
                ),
                "document_changes": {
                    "updated_html": self.render_html(session_id, document.id),
                    "version_id": f"v{session.version}",
                    "changes_summary": f"{len(applied)} applied, {denied} denied",
                    "document_id": document.id,
                },
                "usage": {
                    "monthly_used": self.ops_used,
                    "monthly_remaining": self.ops_limit - self.ops_used,
                },
            }
            return {"status": "ok", "applied": len(applied)}
        if feedback_items and job.round < 2:
            job.round += 1
            revision_note = "; ".join(str(c["feedback"])[:80] for c in feedback_items)
            job.message = f"{job.message} (revision note: {revision_note})"
            job.proposals = self._draft_proposals(job.message, document, round=job.round)
            job.awaiting_kind = None
            return {"status": "ok", "revised": True, "round": job.round}
        job.status = STATUS_COMPLETED
        job.awaiting_kind = None
        job.result = {
            "response": "All proposed changes were rejected; the document was left unchanged.",
            "document_changes": {
                "updated_html": self.render_html(session_id, document.id),
                "version_id": f"v{session.version}",
                "changes_summary": "no changes applied",
            },
            "usage": {"monthly_used": self.ops_used},
        }
        return {"status": "ok", "applied": 0}

    def _continue(self, session_id: str, payload: dict[str, Any]) -> dict[str, Any]:
        job = self.jobs.get(str(payload.get("job_id")))
        if job is None:
            raise SuperDocsAPIError(404, "job_not_found", "unknown job")
        if job.status != STATUS_AWAITING_APPROVAL or job.awaiting_kind != AWAITING_KIND_CONTINUE:
            raise SuperDocsAPIError(
                409, "not_paused_for_continue", "job is not paused at a continue prompt"
            )
        if payload.get("continue", False):
            job.awaiting_kind = None
            job.status = STATUS_COMPLETED
            session = self.sessions[job.session_id]
            job.result = {
                "response": "Large edit finished.",
                "document_changes": {
                    "updated_html": self.render_html(job.session_id),
                    "version_id": f"v{session.version}",
                    "changes_summary": "continued edit finished",
                },
                "usage": {"monthly_used": self.ops_used},
            }
        else:
            job.status = STATUS_COMPLETED
            job.awaiting_kind = None
            job.result = {
                "response": "Stopped at your request; partial work kept.",
                "document_changes": {},
            }
        return {"status": "ok"}

    def _cancel(self, job_id: str) -> dict[str, Any]:
        job = self.jobs.get(job_id)
        if job is None:
            raise SuperDocsAPIError(404, "job_not_found", f"unknown job {job_id!r}")
        if job.status in (STATUS_PENDING, STATUS_IN_PROGRESS):
            job.status = "cancelled"
            return {"status": "cancelled"}
        raise SuperDocsAPIError(
            409, "not_cancellable", f"job in status {job.status!r} cannot be cancelled"
        )

    def _list_session_jobs(self, session_id: str) -> dict[str, Any]:
        jobs = [
            {"job_id": j.id, "status": j.status, "metadata": {"awaiting_kind": j.awaiting_kind}}
            for j in sorted(self.jobs.values(), key=lambda x: x.id)
            if j.session_id == session_id
        ]
        return {"jobs": jobs}

    def _export(self, payload: dict[str, Any]) -> dict[str, Any]:
        session_id = payload.get("session_id")
        fmt = payload.get("format", "docx")
        if not session_id or session_id not in self.sessions:
            raise SuperDocsAPIError(404, "document_not_found", f"unknown session {session_id!r}")
        session = self.sessions[session_id]
        options = payload.get("options") or {}
        document = session.document_or_focused(payload.get("document_id"))
        filename = options.get("filename") or document.title.rsplit(".", 1)[0]
        if fmt == "docx":
            content = self._render_docx(document)
            out_name = f"{filename}.docx"
        elif fmt == "markdown":
            content = "\n\n".join(b["text"] for b in document.blocks).encode("utf-8")
            out_name = f"{filename}.md"
        else:
            content = self.render_html(session_id, document.id).encode("utf-8")
            out_name = f"{filename}.{fmt}"
        return {
            "filename": out_name,
            "format": fmt,
            "content_b64": base64.b64encode(content).decode("ascii"),
            "version_id": f"v{session.version}",
            "document_id": document.id,
        }

    def _render_docx(self, document: FakeDocument) -> bytes:
        import docx

        out = docx.Document()
        for block in document.blocks:
            tag = block["tag"]
            if tag.startswith("h"):
                out.add_heading(block["text"], level=int(tag[1]))
            else:
                out.add_paragraph(block["text"])
        buffer = io.BytesIO()
        out.save(buffer)
        return buffer.getvalue()

    def render_html(self, session_id: str, document_id: str | None = None) -> str:
        session = self.sessions[session_id]
        document = session.document_or_focused(document_id)
        parts = []
        for block in document.blocks:
            parts.append(
                f'<{block["tag"]} data-chunk-id="{block["chunk_id"]}">'
                f'{html_mod.escape(block["text"])}</{block["tag"]}>'
            )
        return "".join(parts)

    def block_texts(self, session_id: str, document_id: str | None = None) -> list[str]:
        session = self.sessions[session_id]
        return [b["text"] for b in session.document_or_focused(document_id).blocks]

    def _draft_proposals(
        self, instruction: str, document: FakeDocument, round: int = 0
    ) -> list[dict[str, Any]]:
        proposals: list[dict[str, Any]] = []
        clauses = [c.strip() for c in re.split(r"[;\n]", instruction.lower()) if c.strip()]
        for clause in clauses:
            self._draft_clause(clause, document, proposals)
        if not proposals:
            fallback = instruction.strip().rstrip(". ")
            proposals.append(
                self._make_proposal(
                    "create",
                    old_text=None,
                    new_text=f"Amendment agreed during negotiation: {fallback}.",
                    explanation="Recorded the negotiation request as a proposed amendment.",
                    target_chunk=None,
                    insert_after=document.blocks[-1]["chunk_id"] if document.blocks else None,
                )
            )
        limited = proposals[:3]
        if round:
            for p in limited:
                p["ai_explanation"] += f" [revised, round {round}]"
        return limited

    def _draft_clause(
        self, clause: str, document: FakeDocument, proposals: list[dict[str, Any]]
    ) -> None:
        cap_pattern = (
            r"cap\s+(?:the\s+)?liability\s+(?:at|to)\s+\$?([\d,.]+)\s*(k|m|thousand|million)?"
        )
        cap = re.search(cap_pattern, clause)
        if cap:
            amount_raw = cap.group(1).replace(",", "")
            unit = cap.group(2)
            multipliers = {"k": "000", "thousand": "000", "m": ",000,000", "million": ",000,000"}
            multiplier = multipliers.get(unit or "", "")
            amount = f"${amount_raw}{multiplier}"
            target = next((b for b in document.blocks if "liability" in b["text"].lower()), None)
            if target:
                sentences = re.split(r"(?<=[.;])\s+", target["text"])
                kept = " ".join(sentences[: max(len(sentences) - 1, 1)]).strip().rstrip(".")
                new_text = (
                    f"{kept}. In no event shall either party's aggregate liability "
                    f"exceed {amount}."
                )
                proposals.append(
                    self._make_proposal(
                        "edit",
                        old_text=target["text"],
                        new_text=new_text,
                        explanation=(
                            f"Capped aggregate liability at {amount}, per the negotiation "
                            "request."
                        ),
                        target_chunk=target["chunk_id"],
                    )
                )
            return

        delete_match = re.search(
            r"\bdelete\b\s+(?:the\s+)?(?:section\s+(?:on|about|covering)\s+)?(.+?)(?:\s+from\b|\s*$)",
            clause,
        )
        if delete_match:
            term = delete_match.group(1).strip().rstrip(".,")
            if term and term != "it":
                target = next((b for b in document.blocks if term in b["text"].lower()), None)
                if target:
                    proposals.append(
                        self._make_proposal(
                            "delete",
                            old_text=target["text"],
                            new_text=None,
                            explanation=f"Removed the passage about {term} as requested.",
                            target_chunk=target["chunk_id"],
                        )
                    )
            return

        add_match = re.search(r"\badd\b[:\s]+(.*?)(?:\s+after\s+(.+?))?\s*$", clause)
        if add_match and add_match.group(1):
            clause_text = add_match.group(1).strip().rstrip(". ")
            anchor = (add_match.group(2) or "").strip()
            after_chunk = None
            if anchor:
                found = next((b for b in document.blocks if anchor in b["text"].lower()), None)
                if found:
                    after_chunk = found["chunk_id"]
            proposals.append(
                self._make_proposal(
                    "create",
                    old_text=None,
                    new_text=f"{clause_text[0].upper()}{clause_text[1:]}.",
                    explanation=f"Added a new clause covering: {clause_text}.",
                    target_chunk=None,
                    insert_after=(
                        after_chunk
                        or (document.blocks[-1]["chunk_id"] if document.blocks else None)
                    ),
                )
            )
            return

        keyword_blocks = sorted(
            (b for b in document.blocks if len(b["text"]) > 40),
            key=lambda b: sum(w in b["text"].lower() for w in clause.split()),
            reverse=True,
        )
        if keyword_blocks:
            overlap = sum(w in keyword_blocks[0]["text"].lower() for w in clause.split())
            if overlap >= 2:
                target = keyword_blocks[0]
                new_text = (
                    f"{target['text'].rstrip('.')}. "
                    f"As agreed during negotiation, this clause is amended per: {clause}."
                )
                proposals.append(
                    self._make_proposal(
                        "edit",
                        old_text=target["text"],
                        new_text=new_text,
                        explanation=f"Adjusted an affected passage per the request: {clause}.",
                        target_chunk=target["chunk_id"],
                    )
                )

    def _make_proposal(
        self,
        operation: str,
        old_text: str | None,
        new_text: str | None,
        explanation: str,
        target_chunk: str | None,
        insert_after: str | None = None,
    ) -> dict[str, Any]:
        proposal: dict[str, Any] = {
            "change_id": self._next_id("ch"),
            "operation": operation,
            "chunk_id": target_chunk if operation != "create" else None,
            "old_html": f"<p>{html_mod.escape(old_text)}</p>" if old_text else None,
            "new_html": f"<p>{html_mod.escape(new_text)}</p>" if new_text else None,
            "ai_explanation": explanation,
            "insert_after_chunk_id": insert_after,
        }
        return proposal

    def _apply(
        self,
        session: FakeSession,
        accepted_ids: list[str],
        proposals: list[dict[str, Any]],
        document: FakeDocument | None = None,
    ) -> list[dict[str, Any]]:
        target = document if document is not None else session.focused
        applied: list[dict[str, Any]] = []
        by_id = {p["change_id"]: p for p in proposals}
        for change_id in accepted_ids:
            proposal = by_id.get(change_id)
            if proposal is None:
                continue
            if proposal["operation"] == "edit":
                for index, block in enumerate(target.blocks):
                    if block["chunk_id"] == proposal["chunk_id"]:
                        target.blocks[index] = {
                            "chunk_id": block["chunk_id"],
                            "tag": block["tag"],
                            "text": html_mod.unescape(
                                re.sub(r"<[^>]+>", "", proposal["new_html"] or "")
                            ),
                        }
                        break
            elif proposal["operation"] == "delete":
                target_id = proposal["chunk_id"]
                target.blocks = [b for b in target.blocks if b["chunk_id"] != target_id]
            elif proposal["operation"] == "create":
                new_block = self._new_chunk(
                    "p", html_mod.unescape(re.sub(r"<[^>]+>", "", proposal["new_html"] or ""))
                )
                insert_at = len(target.blocks)
                if proposal.get("insert_after_chunk_id"):
                    for index, block in enumerate(target.blocks):
                        if block["chunk_id"] == proposal["insert_after_chunk_id"]:
                            insert_at = index + 1
                            break
                target.blocks.insert(insert_at, new_block)
            applied.append(proposal)
        if applied:
            session.version += 1
        return applied
