from __future__ import annotations

import html as html_mod
import io
import re


def extract_document_text(filename: str, raw: bytes) -> str:
    """Best-effort plain text for search indexing.

    Never raises: a document we cannot parse simply contributes no searchable text
    rather than breaking the whole search. Returning "" is honest -- the caller reports
    no hits for that file instead of guessing at its contents.
    """
    lower = filename.lower()
    try:
        if lower.endswith(".docx"):
            import docx

            document = docx.Document(io.BytesIO(raw))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells if cell.text.strip())
            return "\n".join(parts)
        if lower.endswith((".html", ".htm")):
            text = raw.decode("utf-8", errors="replace")
            return html_mod.unescape(re.sub(r"<[^>]+>", " ", text)).strip()
        return raw.decode("utf-8", errors="replace")
    except Exception:
        return ""
