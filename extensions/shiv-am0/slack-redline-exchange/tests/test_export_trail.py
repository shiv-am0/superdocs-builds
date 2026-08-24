from __future__ import annotations

import io
import zipfile

import docx

from redline.core.store import ProposalRow
from redline.export.track_changes import build_redline_docx
from tests.conftest import build_harness, make_contract_docx

NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _revision_authors(content: bytes, tag: str) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml = archive.read("word/document.xml")
    from lxml import etree

    tree = etree.fromstring(xml)
    return [el.get(f"{NS}author") for el in tree.iter(f"{NS}{tag}")]


async def test_export_trail_marks_committed_changes_with_track_changes(tmp_path):
    h = build_harness(tmp_path, approval_policy="dual_consent")
    deal = await h.sim.send_file_share("vendor", "MSA.docx", make_contract_docx())
    deal_id = deal["deal_id"]

    # a mixed batch: one proposal approved by both sides, one rejected by at least one --
    # since at least one change in the job is accepted, SuperDocs applies the approved
    # change and discards the rejected one directly, rather than looping into a revision
    # round (that only happens when *every* change in a job's batch is rejected).
    instruction = "Delete the section on Term; add audit rights after liability"
    job = await h.sim.run_command("customer", f"propose {deal_id} {instruction}")
    proposals = h.store.proposals_for_job(job["job_id"])
    assert len(proposals) >= 2
    keep, drop = proposals[0], proposals[1]
    await h.service.decide(deal_id, keep.id, "vendor", "U_VENDOR", True)
    await h.service.decide(deal_id, keep.id, "customer", "U_CUSTOMER", True)
    await h.service.decide(
        deal_id, drop.id, "vendor", "U_VENDOR", False, feedback="We need the term clause kept"
    )
    resolved_drop = await h.service.decide(deal_id, drop.id, "customer", "U_CUSTOMER", False)
    assert resolved_drop.state == "rejected"

    exported = await h.service.export_redline(deal_id)
    assert exported.filename.endswith("_redline.docx")

    ins_authors = _revision_authors(exported.content, "ins")
    del_authors = _revision_authors(exported.content, "del")
    assert "Customer" in ins_authors + del_authors, "the committed change was proposed by Customer"

    out_doc = docx.Document(io.BytesIO(exported.content))
    full_text = "\n".join(p.text for p in out_doc.paragraphs)
    assert "Appendix: Rejected Proposals" in full_text
    assert "We need the term clause kept" in _appendix_table_text(out_doc)
    assert "Unmatched Changes" not in full_text


def _appendix_table_text(document: docx.Document) -> str:
    chunks = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def _make_proposal(**overrides) -> ProposalRow:
    base = dict(
        id="prop_1",
        deal_id="deal_1",
        job_id="job_1",
        change_id="ch_1",
        operation="edit",
        old_html="<p>text that is definitely not in the source document</p>",
        new_html="<p>new text</p>",
        ai_explanation="test",
        proposed_by_side="vendor",
        state="committed",
    )
    base.update(overrides)
    return ProposalRow(**base)


def test_unmatched_committed_change_is_reported_not_fabricated():
    original = make_contract_docx()
    proposal = _make_proposal()
    content = build_redline_docx("MSA.docx", original, [proposal])
    out_doc = docx.Document(io.BytesIO(content))
    full_text = "\n".join(p.text for p in out_doc.paragraphs)
    assert "Unmatched Changes" in full_text
    ins_authors = _revision_authors(content, "ins")
    assert "Vendor" not in ins_authors, "an unmatched change must never be silently placed inline"


def test_non_docx_original_falls_back_honestly():
    proposal = _make_proposal(operation="delete", old_html="<p>irrelevant</p>", new_html=None)
    content = build_redline_docx("MSA.md", b"# Heading\n\nSome body text.", [proposal])
    out_doc = docx.Document(io.BytesIO(content))
    full_text = "\n".join(p.text for p in out_doc.paragraphs)
    assert "no inline redlines" in full_text.lower()
    assert "Unmatched Changes" in full_text
