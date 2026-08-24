from __future__ import annotations

import hashlib

import pytest

from redline.core.documents import extract_document_text
from redline.core.negotiation import NegotiationError
from tests.conftest import CAP_INSTRUCTION, build_harness, make_contract_docx


def make_amendment_docx() -> bytes:
    import io

    import docx

    document = docx.Document()
    document.add_heading("Amendment No. 1", level=1)
    document.add_paragraph(
        "This Amendment modifies the Master Services Agreement. The parties agree that "
        "invoices shall be settled within 45 days of receipt."
    )
    document.add_paragraph(
        "Section 9: Audit Rights. Customer may audit Vendor's records once per year."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_supporting_document_joins_deal_without_touching_the_contract(tmp_path):
    h = build_harness(tmp_path)
    deal = await h.sim.send_file_share("vendor", "MSA.docx", make_contract_docx())
    deal_id = deal["deal_id"]

    contract_export_before = await h.service.export_clean(deal_id)

    result = await h.sim.send_file_share(
        "customer",
        "Amendment_1.docx",
        make_amendment_docx(),
        attach_to_deal_id=deal_id,
    )
    assert result["deal_id"] == deal_id

    documents = h.service.list_documents(deal_id)
    assert [d.filename for d in documents] == ["MSA.docx", "Amendment_1.docx"]
    assert [d.role for d in documents] == ["contract", "supporting"]

    # the authoritative contract is unchanged: adding a document altered no text.
    # Compared as extracted text, not raw bytes -- a .docx is a ZIP and ZIP headers
    # embed a modification timestamp, so two exports taken a second apart differ in
    # bytes while being identical documents.
    contract_export_after = await h.service.export_clean(deal_id)
    assert extract_document_text(
        contract_export_after.filename, contract_export_after.content
    ) == extract_document_text(
        contract_export_before.filename, contract_export_before.content
    )


async def test_untargeted_propose_still_edits_the_contract(tmp_path):
    h = build_harness(tmp_path)
    deal = await h.sim.send_file_share("vendor", "MSA.docx", make_contract_docx())
    deal_id = deal["deal_id"]
    await h.sim.send_file_share(
        "customer", "Amendment_1.docx", make_amendment_docx(), attach_to_deal_id=deal_id
    )

    job = await h.sim.run_command("vendor", f"propose {deal_id} {CAP_INSTRUCTION}")
    proposals = h.store.proposals_for_job(job["job_id"])
    assert proposals

    contract = next(d for d in h.service.list_documents(deal_id) if d.role == "contract")
    assert all(p.document_id == contract.id for p in proposals), (
        "an untargeted propose must edit the contract, not the most recently uploaded file"
    )


async def test_propose_can_target_a_supporting_document_by_name(tmp_path):
    h = build_harness(tmp_path)
    deal = await h.sim.send_file_share("vendor", "MSA.docx", make_contract_docx())
    deal_id = deal["deal_id"]
    await h.sim.send_file_share(
        "customer", "Amendment_1.docx", make_amendment_docx(), attach_to_deal_id=deal_id
    )
    amendment = next(
        d for d in h.service.list_documents(deal_id) if d.filename == "Amendment_1.docx"
    )

    job = await h.sim.run_command(
        "customer", f"propose {deal_id} --doc Amendment_1 Change the payment window to 30 days"
    )
    proposals = h.store.proposals_for_job(job["job_id"])
    assert proposals
    assert all(p.document_id == amendment.id for p in proposals)


async def test_unknown_document_target_is_refused_with_a_helpful_message(tmp_path):
    h = build_harness(tmp_path)
    deal = await h.sim.send_file_share("vendor", "MSA.docx", make_contract_docx())
    deal_id = deal["deal_id"]

    with pytest.raises(NegotiationError) as excinfo:
        await h.service.propose(
            deal_id,
            "vendor",
            "U_VENDOR",
            CAP_INSTRUCTION,
            source_channel_id=h.settings.shared_channel,
            document_id="does_not_exist.docx",
        )
    message = str(excinfo.value)
    assert "does_not_exist.docx" in message
    assert "MSA.docx" in message, "the error should name the documents that DO exist"


async def test_documents_must_be_shared_in_the_shared_channel(tmp_path):
    h = build_harness(tmp_path)
    deal = await h.sim.send_file_share("vendor", "MSA.docx", make_contract_docx())
    deal_id = deal["deal_id"]

    from redline.core import boundary

    with pytest.raises(boundary.BoundaryViolation):
        await h.service.add_document(
            deal_id,
            "vendor",
            "U_VENDOR",
            "Secret_Strategy.docx",
            make_amendment_docx(),
            source_channel_id=h.settings.vendor_internal_channel,
        )
    assert len(h.service.list_documents(deal_id)) == 1


async def test_refresh_redraws_a_pending_card_that_was_clobbered(tmp_path):
    """Recovery for a card that no longer shows what the database says.

    The store is authoritative; a Slack card is just a remote copy of it. Redrawing is
    therefore always safe and never invents a decision -- it only makes Slack agree with
    what was already recorded.
    """
    h = build_harness(tmp_path)
    deal = await h.sim.send_file_share("vendor", "MSA.docx", make_contract_docx())
    deal_id = deal["deal_id"]
    job = await h.sim.run_command("vendor", f"propose {deal_id} {CAP_INSTRUCTION}")
    proposal = h.store.proposals_for_job(job["job_id"])[0]

    # something overwrites the card in place, exactly as a bad respond() did
    await h.service.messenger.update_message(
        h.settings.shared_channel, proposal.card_ts, [], "clobbered"
    )

    redrawn = await h.service.refresh_cards(deal_id)
    assert redrawn == 1
    restored = [
        m
        for m in h.service.messenger.channels[h.settings.shared_channel]
        if m.ts == proposal.card_ts
    ][0]
    assert "clobbered" not in restored.text
    assert h.store.get_proposal(proposal.id).state == "pending", "redraw must not decide"


async def test_refresh_leaves_already_decided_proposals_alone(tmp_path):
    """Only still-pending cards are redrawn; settled ones keep their resolved state."""
    h = build_harness(tmp_path)
    deal = await h.sim.send_file_share("vendor", "MSA.docx", make_contract_docx())
    deal_id = deal["deal_id"]
    job = await h.sim.run_command("vendor", f"propose {deal_id} {CAP_INSTRUCTION}")
    proposal = h.store.proposals_for_job(job["job_id"])[0]
    await h.sim.click("vendor", proposal.id, "approve")
    await h.sim.click("customer", proposal.id, "approve")

    assert await h.service.refresh_cards(deal_id) == 0


async def test_attaching_a_document_leaves_the_contract_byte_identical(tmp_path):
    """The card promises the contract is untouched, so prove it at the byte level.

    The sibling test compares exported text, because a .docx is a ZIP and its headers
    carry a timestamp that changes between two exports of identical content. The bytes
    we actually store are not subject to that, so this is where the promise is provable
    rather than merely asserted.
    """
    h = build_harness(tmp_path)
    deal = await h.sim.send_file_share("vendor", "MSA.docx", make_contract_docx())
    deal_id = deal["deal_id"]
    contract = next(d for d in h.service.list_documents(deal_id) if d.role == "contract")
    before = h.store.get_document(contract.id).original_bytes
    digest_before = hashlib.sha256(before).hexdigest()

    await h.sim.send_file_share(
        "customer", "Amendment_1.docx", make_amendment_docx(), attach_to_deal_id=deal_id
    )
    await h.sim.run_command(
        "customer", f"propose {deal_id} --doc Amendment_1 Change the payment window to 30 days"
    )

    after = h.store.get_document(contract.id).original_bytes
    assert hashlib.sha256(after).hexdigest() == digest_before
    assert after == before, "attaching or editing another document must not touch the contract"
