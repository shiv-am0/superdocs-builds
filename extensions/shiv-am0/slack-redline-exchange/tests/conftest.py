from __future__ import annotations

import io

import docx
import pytest

from redline.config import Settings
from redline.core.negotiation import NegotiationService
from redline.core.store import Store
from redline.slack.simulator import InMemoryMessenger, TwoWorkspaceSimulator
from redline.superdocs.client import SuperDocsClient
from redline.superdocs.fake import FakeSuperDocs

VENDOR_TEAM = "T_VENDOR"
CUSTOMER_TEAM = "T_CUSTOMER"
VENDOR_INTERNAL = "C_INT_VENDOR"
CUSTOMER_INTERNAL = "C_INT_CUSTOMER"
SHARED_CHANNEL = "C_SHARED"

CAP_INSTRUCTION = "Cap liability at $50k in the liability section"


def make_contract_docx(extra_paragraph: str | None = None) -> bytes:
    document = docx.Document()
    document.add_heading("Mutual Non-Disclosure Agreement", level=1)
    document.add_paragraph(
        'This Agreement is made between Acme Corp ("Vendor") and Globex Ltd ("Customer").'
    )
    document.add_paragraph(
        "Section 1: Confidential Information. Each party may disclose business information."
    )
    document.add_paragraph(
        "Section 5: Term. This Agreement runs for two years from the Effective Date."
    )
    document.add_paragraph(
        "Section 8: Liability. Each party's aggregate liability arising out of this "
        "Agreement is unlimited and without cap of any kind."
    )
    if extra_paragraph:
        document.add_paragraph(extra_paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_settings(tmp_path, **overrides) -> Settings:
    defaults: dict = dict(
        fake_mode=True,
        database_path=str(tmp_path / "redline.db"),
        approval_policy="dual_consent",
        poll_interval_seconds=0.0,
        warn_after_seconds=9999,
        max_job_wait_seconds=30,
        vendor_team_id=VENDOR_TEAM,
        customer_team_id=CUSTOMER_TEAM,
        vendor_internal_channel=VENDOR_INTERNAL,
        customer_internal_channel=CUSTOMER_INTERNAL,
        shared_channel=SHARED_CHANNEL,
    )
    defaults.update(overrides)
    return Settings(**defaults)


class Harness:
    def __init__(
        self,
        settings: Settings,
        store: Store,
        fake: FakeSuperDocs,
        service: NegotiationService,
    ):
        self.settings = settings
        self.store = store
        self.fake = fake
        self.service = service
        self.sim = TwoWorkspaceSimulator(service=service, settings=settings)

    @property
    def messenger(self) -> InMemoryMessenger:
        return self.sim.messenger


def build_harness(tmp_path, **setting_overrides) -> Harness:
    settings = make_settings(tmp_path, **setting_overrides)
    store = Store(settings.database_path)
    fake = FakeSuperDocs()
    client = SuperDocsClient(fake)
    messenger = InMemoryMessenger()
    service = NegotiationService(settings, store, client, messenger)
    return Harness(settings, store, fake, service)


@pytest.fixture
def harness(tmp_path):
    return build_harness(tmp_path)
