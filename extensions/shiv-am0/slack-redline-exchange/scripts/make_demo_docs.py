from __future__ import annotations

import io
from pathlib import Path

import docx

DEMO_DIR = Path(__file__).resolve().parent.parent / "demo"
OUTPUT_PATH = DEMO_DIR / "Acme_Globex_MSA.docx"
AMENDMENT_PATH = DEMO_DIR / "Acme_Globex_Amendment_1.docx"


def build_msa() -> bytes:
    document = docx.Document()
    document.add_heading("Master Services Agreement", level=1)
    document.add_paragraph(
        'This Master Services Agreement ("Agreement") is entered into between Acme Corp '
        '("Vendor") and Globex Ltd ("Customer"), collectively the "Parties". Both are '
        "fictional companies invented for this demonstration."
    )
    document.add_heading("1. Term", level=2)
    document.add_paragraph(
        "This Agreement commences on the Effective Date and continues for an initial "
        "term of two years, renewing automatically for successive one-year terms unless "
        "either Party gives 60 days' written notice of non-renewal."
    )
    document.add_heading("2. Confidentiality", level=2)
    document.add_paragraph(
        "Each Party may disclose confidential business, technical, and financial "
        "information to the other Party in connection with this Agreement, and the "
        "receiving Party shall protect that information with at least reasonable care."
    )
    document.add_heading("3. Limitation of Liability", level=2)
    document.add_paragraph(
        "Each Party's aggregate liability arising out of or related to this Agreement "
        "is unlimited and without cap of any kind, whether in contract, tort, or otherwise."
    )
    document.add_heading("4. Termination", level=2)
    document.add_paragraph(
        "Either Party may terminate this Agreement for material breach if the breach "
        "is not cured within 30 days of written notice."
    )
    document.add_heading("5. Governing Law", level=2)
    document.add_paragraph(
        "This Agreement is governed by the laws of the State of Delaware, without "
        "regard to its conflict of laws principles."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_amendment() -> bytes:
    document = docx.Document()
    document.add_heading("Amendment No. 1 to the Master Services Agreement", level=1)
    document.add_paragraph(
        "This Amendment is entered into between Acme Corp and Globex Ltd and modifies "
        "the Master Services Agreement dated the Effective Date. Both are fictional "
        "companies invented for this demonstration."
    )
    document.add_heading("A1. Payment Terms", level=2)
    document.add_paragraph(
        "Invoices shall be settled within 45 days of receipt. Late payments accrue "
        "interest at 1% per month."
    )
    document.add_heading("A2. Audit Rights", level=2)
    document.add_paragraph(
        "Customer may audit Vendor's records relating to this Agreement no more than "
        "once per calendar year, on 30 days' written notice."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def main() -> None:
    DEMO_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_bytes(build_msa())
    print(f"wrote {OUTPUT_PATH}")
    AMENDMENT_PATH.write_bytes(build_amendment())
    print(f"wrote {AMENDMENT_PATH}")


if __name__ == "__main__":
    main()
